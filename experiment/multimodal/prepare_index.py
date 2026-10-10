"""Extract fixed, gold-blind chunks from the 14 input assets only."""

from __future__ import annotations

import hashlib
import json
import math
import subprocess
import time
from pathlib import Path
from typing import Any

import fitz
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation


INPUT_ROOT = Path("/benchmark/inputs")
SOURCE_MANIFEST = Path("/benchmark/source-manifest.json")
RUN_ROOT = Path("/benchmark/run")
DERIVED_ROOT = RUN_ROOT / "derived"
INDEX_PATH = RUN_ROOT / "index.jsonl"
HASH_PATH = RUN_ROOT / "input-hashes.json"
PREP_PATH = RUN_ROOT / "preparation-summary.json"
REVISION = "914f7f89142e33e77833254d9c9b90c3cef7303b"
AUDIO_CHUNK_SECONDS = 15.0
VIDEO_CHUNK_SECONDS = 6.0
VIDEO_PROCESSING = {
    "fps": 1,
    "max_frames": 32,
    "overflow_strategy": "uniform",
    "add_timestamps": False,
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def emit(chunks: list[dict[str, Any]], asset_id: str, source_path: str, modality: str, locator: str,
         text: str | None = None, title: str | None = None, media_path: str | None = None,
         interval: dict[str, float] | None = None, units: list[str] | None = None) -> None:
    if not (text and text.strip()) and not media_path:
        return
    chunks.append({
        "chunk_id": f"{asset_id}:{len(chunks) + 1:04d}",
        "asset_id": asset_id,
        "source_path": source_path,
        "modality": modality,
        "locator": locator,
        "evidence_units": units or [],
        "title": title,
        "text": text.strip() if text else None,
        "media_path": media_path,
        "interval_seconds": interval,
    })


def docx_chunks(path: Path, asset_id: str, source_path: str, chunks: list[dict[str, Any]]) -> None:
    document = Document(path)
    nodes = list(document.element.body.iterchildren())
    source_title = ""
    for node in nodes:
        if node.tag != qn("w:p"):
            continue
        paragraph = Paragraph(node, document)
        style_name = getattr(paragraph.style, "name", "") or ""
        candidate = (paragraph.text or "").strip()
        if candidate and style_name.lower().startswith("heading"):
            source_title = candidate
            break
    source_title = source_title or "Internal document"
    section_stack: list[tuple[int, str]] = []
    current: list[str] = []
    current_heading = "document opening"
    ordinal = 0

    def flush() -> None:
        nonlocal ordinal, current
        body = "\n".join(item for item in current if item.strip()).strip()
        if body:
            ordinal += 1
            locator = " > ".join(name for _, name in section_stack) if section_stack else current_heading
            emit(chunks, asset_id, source_path, "document", f"section:{locator or ordinal}", text=body, title=source_title,
                 units=[f"section:{locator or ordinal}"])
        current = []

    for node in nodes:
        if node.tag == qn("w:p"):
            paragraph = Paragraph(node, document)
            text = (paragraph.text or "").strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                flush()
                try:
                    level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                while section_stack and section_stack[-1][0] >= level:
                    section_stack.pop()
                section_stack.append((level, text))
                current_heading = text
            current.append(text)
        elif node.tag == qn("w:tbl"):
            flush()
            table = Table(node, document)
            ordinal += 1
            row_lines: list[str] = []
            for row_no, row in enumerate(table.rows, start=1):
                cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
                row_lines.append(" | ".join(value for value in cells if value))
            body = "\n".join(row_lines).strip()
            if body:
                locator = f"table:{ordinal}"
                emit(chunks, asset_id, source_path, "document", locator, text=body, title=source_title,
                     units=[locator])
    flush()


def pdf_chunks(path: Path, asset_id: str, source_path: str, chunks: list[dict[str, Any]]) -> None:
    with fitz.open(path) as pdf:
        first_page_text = pdf[0].get_text("text").strip() if len(pdf) else ""
        source_title = next((line.strip() for line in first_page_text.splitlines() if line.strip()), "Internal report")
        for page_no, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            locator = f"page:{page_no}"
            emit(chunks, asset_id, source_path, "pdf", locator, text=text, title=source_title, units=[locator])


def cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xlsx_chunks(path: Path, asset_id: str, source_path: str, chunks: list[dict[str, Any]]) -> None:
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached_book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        for sheet_no, formula_sheet in enumerate(formula_book.worksheets, start=1):
            cached_sheet = cached_book[formula_sheet.title]
            rows = []
            for row_no in range(1, formula_sheet.max_row + 1):
                populated = []
                for col_no in range(1, formula_sheet.max_column + 1):
                    cell = formula_sheet.cell(row_no, col_no)
                    value = cell.value
                    if value is not None and str(value).strip():
                        populated.append((cell, cached_sheet.cell(row_no, col_no).value))
                if populated:
                    rows.append((row_no, populated))
            if not rows:
                continue

            header_row, header_values = rows[0]
            source_title = " | ".join(cell_value(cell.value) for cell, _ in header_values if cell_value(cell.value))
            source_title = source_title or f"Spreadsheet sheet {sheet_no}"
            headers: dict[int, str] = {}
            for cell, _ in header_values:
                value = cell_value(cell.value)
                headers[cell.column] = value or f"column_{cell.column_letter}"

            preamble = ""
            if header_row > 1:
                preamble_rows = []
                for row_no in range(1, header_row):
                    values = [cell_value(formula_sheet.cell(row_no, col_no).value)
                              for col_no in range(1, formula_sheet.max_column + 1)]
                    values = [value for value in values if value]
                    if values:
                        preamble_rows.append(" | ".join(values))
                preamble = "\n".join(preamble_rows)

            for row_no, populated in rows[1:]:
                fields = []
                units = []
                start_col = min(cell.column for cell, _ in populated)
                end_col = max(cell.column for cell, _ in populated)
                for cell, cached_value in populated:
                    header = headers.get(cell.column, f"column_{cell.column_letter}")
                    raw = cell.value
                    if isinstance(raw, str) and raw.startswith("="):
                        value = f"formula {raw}; cached value {cell_value(cached_value) or 'unavailable'}"
                    else:
                        value = cell_value(raw)
                    fields.append(f"{header} = {value}")
                    units.append(cell.coordinate)
                locator = f"sheet:{sheet_no}:rows:{row_no}-{row_no}:cells:{formula_sheet.cell(row_no, start_col).coordinate}:{formula_sheet.cell(row_no, end_col).coordinate}"
                body = " | ".join(fields)
                if preamble and row_no == rows[1][0]:
                    body = f"{preamble}\nColumn headers: " + " | ".join(
                        headers.get(col_no, f"column_{col_no}")
                        for col_no in sorted(headers)
                    ) + "\n" + body
                emit(chunks, asset_id, source_path, "spreadsheet", locator, text=body, title=source_title, units=units)
    finally:
        formula_book.close()
        cached_book.close()


def pptx_chunks(path: Path, asset_id: str, source_path: str, chunks: list[dict[str, Any]]) -> None:
    presentation = Presentation(path)
    for slide_no, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape_no, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                for series in chart.series:
                    values = ", ".join(cell_value(value) for value in series.values)
                    parts.append(f"Chart series {series.name}: {values}")
        locator = f"slide:{slide_no}"
        source_title = next((line.strip() for part in parts for line in part.splitlines() if line.strip()), f"Presentation slide {slide_no}")
        emit(chunks, asset_id, source_path, "presentation", locator,
             text="\n".join(parts), title=source_title, units=[locator])


def ffprobe_duration(path: Path) -> float:
    proc = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)],
        check=True, capture_output=True, text=True,
    )
    return float(json.loads(proc.stdout)["format"]["duration"])


def media_chunks(path: Path, asset: dict[str, Any], chunks: list[dict[str, Any]]) -> dict[str, Any]:
    asset_id = asset["asset_id"]
    extension = asset["format"]
    rel = Path(asset["path"]).relative_to("inputs")
    source = INPUT_ROOT / rel
    duration = ffprobe_duration(source)
    is_audio = extension == "wav"
    step = AUDIO_CHUNK_SECONDS if is_audio else VIDEO_CHUNK_SECONDS
    out_dir = DERIVED_ROOT / ("audio" if is_audio else "video") / asset_id
    out_dir.mkdir(parents=True, exist_ok=True)
    count = math.ceil(duration / step)
    merged_short_audio_tail = False
    if is_audio and count > 1 and duration - (count - 1) * step < 1.0:
        count -= 1
        merged_short_audio_tail = True
    created = []
    for ordinal in range(count):
        start = ordinal * step
        end = duration if ordinal == count - 1 else min(duration, (ordinal + 1) * step)
        span = end - start
        suffix = "wav" if is_audio else "mp4"
        out_path = out_dir / f"segment-{ordinal + 1:02d}.{suffix}"
        if is_audio:
            command = [
                "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
                "-ss", f"{start:.6f}", "-i", str(source), "-t", f"{span:.6f}",
                "-map", "0:a:0", "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le",
                str(out_path),
            ]
        else:
            command = [
                "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
                "-ss", f"{start:.6f}", "-i", str(source), "-t", f"{span:.6f}",
                "-map", "0:v:0", "-an", "-c:v", "libx264", "-preset", "ultrafast",
                "-crf", "0", "-pix_fmt", "yuv420p", "-movflags", "+faststart",
                str(out_path),
            ]
        proc = subprocess.run(command, check=True, capture_output=True, text=True)
        chunk_no = len(chunks) + 1
        locator = f"time:{start:.3f}-{end:.3f}s"
        modality = "audio" if is_audio else "video"
        emit(chunks, asset_id, asset["path"], modality, locator,
             media_path=str(out_path.relative_to(RUN_ROOT)),
             interval={"start": start, "end": end},
             units=[locator])
        created.append({
            "path": str(out_path.relative_to(RUN_ROOT)),
            "sha256": sha256(out_path),
            "bytes": out_path.stat().st_size,
            "start_seconds": start,
            "end_seconds": end,
            "ffmpeg_stderr": proc.stderr,
        })
    return {
        "asset_id": asset_id,
        "source_duration_seconds": duration,
        "chunk_seconds": step,
        "short_audio_tail_merged_into_previous_interval": merged_short_audio_tail,
        "segment_count": len(created),
        "segments": created,
    }


def main() -> None:
    started = time.perf_counter()
    DERIVED_ROOT.mkdir(parents=True, exist_ok=True)
    source_manifest = json.loads(SOURCE_MANIFEST.read_text(encoding="utf-8"))
    assets = source_manifest["assets"]
    if len(assets) != 14:
        raise RuntimeError(f"expected exactly 14 allowlisted assets; found {len(assets)}")

    allowed = {item["path"] for item in assets}
    present = {
        "inputs/" + path.relative_to(INPUT_ROOT).as_posix()
        for path in INPUT_ROOT.rglob("*") if path.is_file()
    }
    if present != allowed:
        raise RuntimeError(f"input allowlist mismatch; unexpected={sorted(present-allowed)} missing={sorted(allowed-present)}")

    input_hashes = []
    for asset in assets:
        rel = Path(asset["path"]).relative_to("inputs")
        path = INPUT_ROOT / rel
        actual = sha256(path)
        if path.stat().st_size != asset["bytes"] or actual != asset["sha256"]:
            raise RuntimeError(f"input hash mismatch for {asset['asset_id']}: expected {asset['sha256']}, got {actual}")
        input_hashes.append({"asset_id": asset["asset_id"], "path": asset["path"], "bytes": path.stat().st_size, "sha256": actual})

    chunks: list[dict[str, Any]] = []
    media = []
    for asset in assets:
        ext = asset["format"]
        path = INPUT_ROOT / Path(asset["path"]).relative_to("inputs")
        if ext == "docx":
            docx_chunks(path, asset["asset_id"], asset["path"], chunks)
        elif ext == "pdf":
            pdf_chunks(path, asset["asset_id"], asset["path"], chunks)
        elif ext == "xlsx":
            xlsx_chunks(path, asset["asset_id"], asset["path"], chunks)
        elif ext == "pptx":
            pptx_chunks(path, asset["asset_id"], asset["path"], chunks)
        elif ext in {"wav", "mp4"}:
            media.append(media_chunks(path, asset, chunks))
        else:
            raise RuntimeError(f"unsupported format in allowlist: {ext}")

    seen_ids: set[str] = set()
    with INDEX_PATH.open("w", encoding="utf-8", newline="\n") as output:
        for chunk in chunks:
            if chunk["chunk_id"] in seen_ids:
                raise RuntimeError(f"duplicate chunk id: {chunk['chunk_id']}")
            seen_ids.add(chunk["chunk_id"])
            output.write(json.dumps(chunk, ensure_ascii=False, sort_keys=True) + "\n")

    write_json(HASH_PATH, {
        "schema": "embeddinggemma2-accuracy-speed-input-hashes/v1",
        "input_count": len(input_hashes),
        "inputs": input_hashes,
        "allowlist_match": True,
        "source_manifest_read": True,
        "gold_or_query_files_read": False,
    })
    write_json(PREP_PATH, {
        "schema": "embeddinggemma2-accuracy-speed-preparation/v1",
        "model_revision": REVISION,
        "chunk_count": len(chunks),
        "chunks_by_asset": {asset["asset_id"]: sum(c["asset_id"] == asset["asset_id"] for c in chunks) for asset in assets},
        "chunks_by_modality": {key: sum(c["modality"] == key for c in chunks) for key in sorted({c["modality"] for c in chunks})},
        "office_rules": {
            "docx": "paragraphs grouped by natural heading section; tables are separate natural chunks",
            "pdf": "one chunk per source page",
            "xlsx": "one chunk per populated data row with column labels and source cell coordinates",
            "pptx": "one chunk per source slide",
        },
        "media_rules": {
            "audio_chunk_seconds": AUDIO_CHUNK_SECONDS,
            "audio_conversion": "ffmpeg to mono 16 kHz PCM WAV, fixed consecutive intervals",
            "audio_model_limit_note": "15-second segments stay below the approximately 30-second processor input limit",
            "audio_tail_rule": "if the final residue is shorter than 1 second, extend the preceding final interval to the source end; never emit a sub-second audio chunk",
            "video_chunk_seconds": VIDEO_CHUNK_SECONDS,
            "video_conversion": "ffmpeg fixed consecutive clips, H.264 CRF 0, yuv420p",
            "video_processor": VIDEO_PROCESSING,
            "media_prompt": "empty prompt; do not prepend the text document task prefix to media-only inputs",
            "no_question_conditioned_cuts": True,
        },
        "media_segments": media,
        "qa_or_source_material_used": False,
        "index_path": str(INDEX_PATH),
        "index_sha256": sha256(INDEX_PATH),
        "extraction_seconds": round(time.perf_counter() - started, 6),
        "chunk_id_uniqueness": len(seen_ids) == len(chunks),
    })
    print(json.dumps({"chunks": len(chunks), "input_count": len(input_hashes), "extraction_seconds": round(time.perf_counter() - started, 6)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
