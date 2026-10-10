"""資料を検索単位（チャンク）に分ける。区切り方は experiment/multimodal/prepare_index.py と同じ。

- Word: 見出しごと（表は別チャンク）
- PDF: 1ページごと
- Excel: データ1行ごと（列見出しとセル番地つき）
- PowerPoint: 1スライドごと
- テキスト/Markdown: ファイル全体で1チャンク（実験にはない形式。手軽に試すために追加）
- 音声: 15秒ごと、モノラル16kHzへ変換
- 動画: 6秒ごと
"""
import json
import math
import subprocess
from pathlib import Path
from typing import Any

import fitz
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

AUDIO_CHUNK_SECONDS = 15.0
VIDEO_CHUNK_SECONDS = 6.0
AUDIO_EXT = {".wav", ".mp3", ".m4a", ".flac"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm"}


def emit(chunks, asset_id, source_path, modality, locator, text=None, title=None, media_path=None, interval=None):
    if not (text and text.strip()) and not media_path:
        return
    chunks.append({
        "chunk_id": f"{asset_id}:{len(chunks) + 1:04d}",
        "asset_id": asset_id,
        "source_path": source_path,
        "modality": modality,
        "locator": locator,
        "title": title,
        "text": text.strip() if text else None,
        "media_path": media_path,
        "interval_seconds": interval,
    })


def docx_chunks(path, asset_id, source_path, chunks):
    document = Document(path)
    nodes = list(document.element.body.iterchildren())
    source_title = ""
    for node in nodes:
        if node.tag != qn("w:p"):
            continue
        paragraph = Paragraph(node, document)
        style_name = getattr(paragraph.style, "name", "") or ""
        candidate = (paragraph.text or "").strip()
        if candidate and style_name.lower().startswith("heading"):
            source_title = candidate
            break
    source_title = source_title or path.stem
    section_stack: list[tuple[int, str]] = []
    current: list[str] = []
    current_heading = "document opening"
    ordinal = 0

    def flush():
        nonlocal ordinal, current
        body = "\n".join(item for item in current if item.strip()).strip()
        if body:
            ordinal += 1
            locator = " > ".join(name for _, name in section_stack) if section_stack else current_heading
            emit(chunks, asset_id, source_path, "document", f"section:{locator or ordinal}", text=body, title=source_title)
        current = []

    for node in nodes:
        if node.tag == qn("w:p"):
            paragraph = Paragraph(node, document)
            text = (paragraph.text or "").strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                flush()
                try:
                    level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                while section_stack and section_stack[-1][0] >= level:
                    section_stack.pop()
                section_stack.append((level, text))
                current_heading = text
            current.append(text)
        elif node.tag == qn("w:tbl"):
            flush()
            table = Table(node, document)
            ordinal += 1
            lines = []
            for row in table.rows:
                cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
                lines.append(" | ".join(value for value in cells if value))
            emit(chunks, asset_id, source_path, "document", f"table:{ordinal}", text="\n".join(lines), title=source_title)
    flush()


def pdf_chunks(path, asset_id, source_path, chunks):
    with fitz.open(path) as pdf:
        first = pdf[0].get_text("text").strip() if len(pdf) else ""
        source_title = next((line.strip() for line in first.splitlines() if line.strip()), path.stem)
        for page_no, page in enumerate(pdf, start=1):
            emit(chunks, asset_id, source_path, "pdf", f"page:{page_no}", text=page.get_text("text"), title=source_title)


def cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xlsx_chunks(path, asset_id, source_path, chunks):
    formula_book = openpyxl.load_workbook(path, data_only=False)
    cached_book = openpyxl.load_workbook(path, data_only=True)
    try:
        for sheet_no, sheet in enumerate(formula_book.worksheets, start=1):
            cached = cached_book[sheet.title]
            rows = []
            for row_no in range(1, sheet.max_row + 1):
                populated = [(sheet.cell(row_no, c), cached.cell(row_no, c).value)
                             for c in range(1, sheet.max_column + 1)
                             if sheet.cell(row_no, c).value is not None and str(sheet.cell(row_no, c).value).strip()]
                if populated:
                    rows.append((row_no, populated))
            if not rows:
                continue
            _, header_values = rows[0]
            headers = {cell.column: cell_value(cell.value) or f"column_{cell.column_letter}" for cell, _ in header_values}
            source_title = " | ".join(headers[c] for c in sorted(headers)) or f"sheet {sheet_no}"
            for row_no, populated in rows[1:]:
                fields = []
                for cell, cached_value in populated:
                    raw = cell.value
                    # 数式セルは、式と計算済みの値の両方を文章にする
                    if isinstance(raw, str) and raw.startswith("="):
                        value = f"formula {raw}; cached value {cell_value(cached_value) or 'unavailable'}"
                    else:
                        value = cell_value(raw)
                    fields.append(f"{headers.get(cell.column, f'column_{cell.column_letter}')} = {value}")
                first, last = populated[0][0].coordinate, populated[-1][0].coordinate
                emit(chunks, asset_id, source_path, "spreadsheet", f"sheet:{sheet_no}:rows:{row_no}-{row_no}:cells:{first}:{last}",
                     text=" | ".join(fields), title=source_title)
    finally:
        formula_book.close()
        cached_book.close()


def pptx_chunks(path, asset_id, source_path, chunks):
    for slide_no, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
        title = next((line.strip() for part in parts for line in part.splitlines() if line.strip()), f"slide {slide_no}")
        emit(chunks, asset_id, source_path, "presentation", f"slide:{slide_no}", text="\n".join(parts), title=title)


def text_chunks(path, asset_id, source_path, chunks):
    emit(chunks, asset_id, source_path, "document", "file", text=path.read_text(encoding="utf-8"), title=path.stem)


def media_chunks(path, asset_id, source_path, chunks, derived_root: Path):
    probe = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)],
                           check=True, capture_output=True, text=True)
    duration = float(json.loads(probe.stdout)["format"]["duration"])
    is_audio = path.suffix.lower() in AUDIO_EXT
    step = AUDIO_CHUNK_SECONDS if is_audio else VIDEO_CHUNK_SECONDS
    out_dir = derived_root / asset_id
    out_dir.mkdir(parents=True, exist_ok=True)
    count = math.ceil(duration / step)
    if is_audio and count > 1 and duration - (count - 1) * step < 1.0:
        count -= 1  # 1秒未満の端数は直前の区間に含める
    for ordinal in range(count):
        start = ordinal * step
        end = duration if ordinal == count - 1 else min(duration, (ordinal + 1) * step)
        out_path = out_dir / f"segment-{ordinal + 1:02d}.{'wav' if is_audio else 'mp4'}"
        head = ["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-ss", f"{start:.6f}", "-i", str(path), "-t", f"{end - start:.6f}"]
        if is_audio:
            tail = ["-map", "0:a:0", "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le"]
        else:
            tail = ["-map", "0:v:0", "-an", "-c:v", "libx264", "-preset", "ultrafast", "-crf", "0", "-pix_fmt", "yuv420p"]
        subprocess.run(head + tail + [str(out_path)], check=True, capture_output=True, text=True)
        emit(chunks, asset_id, source_path, "audio" if is_audio else "video", f"time:{start:.3f}-{end:.3f}s",
             media_path=str(out_path), interval={"start": start, "end": end})


def chunk_folder(input_root: Path, derived_root: Path) -> list[dict]:
    """フォルダ内の対応ファイルをすべてチャンクにする。asset_id はファイル名（拡張子なし）。"""
    handlers = {".docx": docx_chunks, ".pdf": pdf_chunks, ".xlsx": xlsx_chunks, ".pptx": pptx_chunks,
                ".txt": text_chunks, ".md": text_chunks}
    chunks: list[dict] = []
    for path in sorted(p for p in input_root.rglob("*") if p.is_file()):
        ext = path.suffix.lower()
        rel = path.relative_to(input_root).as_posix()
        if ext in handlers:
            handlers[ext](path, path.stem, rel, chunks)
        elif ext in AUDIO_EXT | VIDEO_EXT:
            media_chunks(path, path.stem, rel, chunks, derived_root)
        elif path.name != ".gitkeep":
            print(f"skip (unsupported): {rel}")
    return chunks
